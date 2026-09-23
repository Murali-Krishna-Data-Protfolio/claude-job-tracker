"""
reclassify_legacy.py — Re-verify rows that were added by the old fail-open
classifier bug, and discard the ones that don't hold up.

Background (see docs/data-workflow.html Figure 2 / the
job-tracker-english-filter-fix note): job_fetcher.py used to catch any
classification error and add the job anyway with a fabricated
confidence=0.5. Every row added that way has English_Reason starting with
"parse error" — that string is the exact fingerprint of an unverified row.
This script finds those rows, re-runs them through the current (fixed,
fail-closed) classifier using their stored Title/Company/Location/
Description_Preview, and:
  - keeps the row (with real confidence + reason) if it now genuinely
    passes as English-speaking
  - discards the row entirely if it fails, isn't confident enough, or the
    classifier errors again — "discard to out of the pipeline"

Always makes a timestamped backup of the workbook before writing anything.
Rewrites the whole Jobs sheet from scratch (see excel_writer.rewrite_jobs_sheet)
rather than deleting rows in place — repeated ws.delete_rows() calls on this
project's openpyxl version don't reliably compact the sheet and can leave
stale blank rows behind.

Usage:
    python reclassify_legacy.py            # dry run — report scope only, no API calls, no writes
    python reclassify_legacy.py --apply    # actually re-classify and rewrite the sheet
"""

import shutil
import sys
from datetime import datetime
from pathlib import Path

# Job titles/companies are full of accented French text (é, è, ç...); some
# Windows terminals default stdout to cp1252, which crashes on encode instead
# of just failing to render the glyph. Never let a print() kill a run mid-way
# through a destructive operation.
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except (AttributeError, ValueError):
    pass

from openpyxl import load_workbook

from config import OUTPUT_PATH
from excel_writer import _build_dashboard, rewrite_jobs_sheet, save_workbook
from job_fetcher import filter_english_jobs

LEGACY_MARKER = "parse error"


def _get_all_rows(ws):
    """Return list of (row_number, row_dict) for every populated row."""
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    out = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[0]:
            out.append((row_num, dict(zip(header, row))))
    return out


def main():
    apply = "--apply" in sys.argv

    if not Path(OUTPUT_PATH).exists():
        print(f"[ERROR] No tracker found at {OUTPUT_PATH}")
        sys.exit(1)

    wb = load_workbook(OUTPUT_PATH)
    ws = wb["Jobs"]
    all_rows = _get_all_rows(ws)
    legacy = [(rn, r) for rn, r in all_rows if LEGACY_MARKER in str(r.get("English_Reason", "")).lower()]

    print(f"Tracker: {OUTPUT_PATH}")
    print(f"Total jobs: {len(all_rows)}")
    print(f"Legacy / unverified rows (English_Reason contains {LEGACY_MARKER!r}): {len(legacy)}")

    if not legacy:
        print("Nothing to reclassify.")
        return

    if not apply:
        print("\nDry run — showing first 10, no API calls made, nothing written:")
        for row_num, r in legacy[:10]:
            print(f"  row {row_num}: {r.get('Title')} @ {r.get('Company')}")
        print("\nRe-run with --apply to actually re-classify and rewrite the sheet.")
        return

    # ── Build job dicts for classification from the stored row data ──────────
    jobs_by_row = {}
    for row_num, r in legacy:
        jobs_by_row[row_num] = {
            "job_id":      r.get("Job_ID", ""),
            "title":       r.get("Title", "") or "",
            "company":     r.get("Company", "") or "",
            "location":    r.get("Location", "") or "",
            "source":      "(legacy — source not stored)",
            "description": r.get("Description_Preview", "") or "",
        }

    print(f"\nRe-classifying {len(jobs_by_row)} rows with the current fail-closed classifier...")
    passing = filter_english_jobs(list(jobs_by_row.values()))
    passing_by_row = {}
    # match back to row numbers by job_id (unique per row)
    id_to_row = {j["job_id"]: rn for rn, j in jobs_by_row.items()}
    for job in passing:
        rn = id_to_row.get(job["job_id"])
        if rn is not None:
            passing_by_row[rn] = job

    kept = sorted(passing_by_row.keys())
    discarded = sorted(rn for rn in jobs_by_row if rn not in passing_by_row)

    print(f"\nKept (genuinely English, re-verified):    {len(kept)}")
    print(f"Discarded (failed / low confidence / error): {len(discarded)}")

    # ── Backup before writing anything ────────────────────────────────────────
    backup_path = Path(OUTPUT_PATH).with_name(
        f"{Path(OUTPUT_PATH).stem}.backup-{datetime.now():%Y%m%d-%H%M%S}.xlsx"
    )
    shutil.copy2(OUTPUT_PATH, backup_path)
    print(f"\nBackup saved: {backup_path}")

    # Build the final row set: every non-legacy row unchanged, plus legacy
    # rows that passed re-verification (with real confidence/reason now),
    # minus legacy rows that were discarded.
    discarded_set = set(discarded)
    keep_rows = []
    for row_num, r in all_rows:
        if row_num in discarded_set:
            continue
        if row_num in passing_by_row:
            job = passing_by_row[row_num]
            r = dict(r)
            r["English_Confidence"] = job.get("english_confidence", "")
            r["English_Reason"] = job.get("english_reason", "")
        keep_rows.append(r)

    rewrite_jobs_sheet(wb, keep_rows)
    _build_dashboard(wb)
    save_workbook(wb)

    print(f"\nDone. {len(kept)} rows re-verified and kept, {len(discarded)} discarded.")
    print(f"Tracker: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
